"""Rehace los archivos de la carpeta entregables/.

1. Copia el Word del artículo (articulo/articulo_cosecha_de_piedras.docx) como
   01_articulo_cosecha_de_piedras.docx y hace su PDF con LibreOffice.
2. Convierte a Word cada documento 02 a 06 (.md), con letra Times New Roman,
   tablas con bordes y cabecera repetida, imágenes ajustadas a la caja y número
   de página al pie.

Uso (desde la raíz del repositorio), después de python3 articulo/generar_docx.py:
    pip install pypandoc_binary python-docx
    python3 entregables/generar_entregables.py
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt

BASE = Path(__file__).resolve().parent
RAIZ = BASE.parent
ARTICULO = RAIZ / "articulo" / "articulo_cosecha_de_piedras.docx"
FONT = "Times New Roman"
MONO = "Courier New"
ANCHO_TEXTO = Inches(8.5) - 2 * Cm(2.5)


def fuente(style, size, bold=None, italic=None, font=FONT):
    style.font.name = font
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), font)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    color = rpr.find(qn("w:color"))
    if color is not None:
        rpr.remove(color)


def espaciado(style, line=1.15, before=0, after=6):
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def plantilla(destino):
    data = subprocess.run(
        [pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True,
    ).stdout
    destino.write_bytes(data)
    doc = Document(str(destino))
    styles = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Definition",
                 "Definition Term", "Table"):
        if name in styles:
            fuente(styles[name], 11)
            espaciado(styles[name])
    if "Compact" in styles:
        espaciado(styles["Compact"], after=2)
    for name, size, before in (("Title", 16, 0), ("Heading 1", 16, 0), ("Heading 2", 13, 14),
                               ("Heading 3", 11.5, 10), ("Heading 4", 11, 8)):
        if name in styles:
            fuente(styles[name], size, bold=True, italic=False)
            espaciado(styles[name], line=1.0, before=before, after=6)
            styles[name].paragraph_format.keep_with_next = True
    if "Block Text" in styles:
        fuente(styles["Block Text"], 10.5, italic=False)
        espaciado(styles["Block Text"], line=1.1, before=4, after=8)
        styles["Block Text"].paragraph_format.left_indent = Cm(1)
        styles["Block Text"].paragraph_format.right_indent = Cm(0.5)
    for name in ("Image Caption", "Caption"):
        if name in styles:
            fuente(styles[name], 9.5, italic=True)
            espaciado(styles[name], line=1.0, before=2, after=10)
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Figure", "Captioned Figure"):
        if name in styles:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            styles[name].paragraph_format.keep_with_next = True
    if "Verbatim Char" in styles:
        fuente(styles["Verbatim Char"], 9.5, font=MONO)
    if "Source Code" in styles:
        fuente(styles["Source Code"], 9.5, font=MONO)
    doc.save(str(destino))


ORDEN = {
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
              "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
              "tblPrChange"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
             "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del",
             "trPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
             "cellDel", "cellMerge", "tcPrChange"],
}


def poner(padre, elemento):
    """Inserta elemento en padre en el orden del esquema de Word, sin duplicarlo."""
    orden = ORDEN[padre.tag.split("}")[1]]
    nombre = elemento.tag.split("}")[1]
    for viejo in padre.findall(qn(f"w:{nombre}")):
        padre.remove(viejo)
    for hijo in padre:
        tag = hijo.tag.split("}")[1]
        if tag in orden and orden.index(tag) > orden.index(nombre):
            hijo.addprevious(elemento)
            return elemento
    padre.append(elemento)
    return elemento


def bordes(tabla):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "8C8C8C")
        b.append(e)
    poner(tblPr, b)
    margen = OxmlElement("w:tblCellMar")
    for lado, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(w))
        e.set(qn("w:type"), "dxa")
        margen.append(e)
    poner(tblPr, margen)


def anchos(tabla):
    """Reparte el ancho de la caja según el largo medio del texto de cada columna.

    Ninguna columna queda más estrecha que su palabra más larga, para que no se
    partan palabras como «Opción» o «Figuras».
    """
    filas = tabla.rows
    n = len(tabla.columns)
    largos = [0.0] * n
    palabra = [0] * n
    for fila in filas:
        for i, celda in enumerate(fila.cells[:n]):
            largos[i] += len(celda.text)
            # Los correos y las direcciones web pueden partirse: no cuentan.
            palabra[i] = max([palabra[i]] + [len(w) for w in celda.text.split()
                                             if "@" not in w and "/" not in w])
    pesos = [max(l / max(len(filas), 1), 4) ** 0.8 for l in largos]
    minimos = [min(max(Cm(0.21) * p + Cm(0.4), Cm(1.2)), int(ANCHO_TEXTO * 0.3)) for p in palabra]
    if sum(minimos) > ANCHO_TEXTO * 0.9:
        minimos = [int(m * ANCHO_TEXTO * 0.9 / sum(minimos)) for m in minimos]
    ws = [ANCHO_TEXTO * p / sum(pesos) for p in pesos]
    fijas = set()
    while True:
        nuevas = {i for i in range(n) if i not in fijas and ws[i] < minimos[i]}
        if not nuevas:
            break
        fijas |= nuevas
        for i in fijas:
            ws[i] = minimos[i]
        libres = [i for i in range(n) if i not in fijas]
        resto = ANCHO_TEXTO - sum(ws[i] for i in fijas)
        suma = sum(pesos[i] for i in libres) or 1
        for i in libres:
            ws[i] = resto * pesos[i] / suma
    ws = [int(w) for w in ws]
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    poner(tblPr, layout)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(Emu(sum(ws)).twips)))
    poner(tblPr, tblW)
    for i, col in enumerate(tabla.columns):
        col.width = Emu(ws[i])
    for fila in filas:
        poner(fila._tr.get_or_add_trPr(), OxmlElement("w:cantSplit"))
        for i, celda in enumerate(fila.cells[:n]):
            celda.width = Emu(ws[i])


def cabecera(tabla):
    fila = tabla.rows[0]
    poner(fila._tr.get_or_add_trPr(), OxmlElement("w:tblHeader"))
    for celda in fila.cells:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "EDEDED")
        poner(celda._tc.get_or_add_tcPr(), shd)
        for par in celda.paragraphs:
            for run in par.runs:
                run.font.bold = True


def numero_de_pagina(seccion):
    par = seccion.footer.paragraphs[0] if seccion.footer.paragraphs else seccion.footer.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = texto
        run._r.append(e)
    run.font.size = Pt(9)
    run.font.name = FONT


def convertir(md, ref):
    docx = md.with_suffix(".docx")
    pypandoc.convert_file(
        str(md), "docx", outputfile=str(docx), format="markdown",
        extra_args=[f"--reference-doc={ref}", f"--resource-path={BASE}",
                    "-M", "lang=es-BO", "--wrap=none"],
    )
    doc = Document(str(docx))
    for seccion in doc.sections:
        seccion.page_width = Inches(8.5)
        seccion.page_height = Inches(11)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(seccion, lado, Cm(2.5))
        # Word exige estos tres atributos de los márgenes; pandoc no los escribe.
        seccion.header_distance = Cm(1.25)
        seccion.footer_distance = Cm(1.25)
        seccion.gutter = Cm(0)
        numero_de_pagina(seccion)
    for tabla in doc.tables:
        bordes(tabla)
        anchos(tabla)
        cabecera(tabla)
        letra = 9 if len(tabla.rows) > 60 else 9.5
        for fila in tabla.rows:
            for celda in fila.cells:
                for par in celda.paragraphs:
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.space_before = Pt(1)
                    for run in par.runs:
                        run.font.size = Pt(letra)
    # Imágenes: dentro de 13 × 14 cm, sin deformarlas.
    for forma in doc.inline_shapes:
        w, h = forma.width, forma.height
        escala = min(Cm(13) / w, Cm(14) / h, 1.0)
        forma.width, forma.height = int(w * escala), int(h * escala)
    props = doc.core_properties
    props.author = "Joaquin Max Cuevas Ventura"
    props.title = md.stem.split("_", 1)[1].replace("_", " ").capitalize()
    doc.save(str(docx))
    return docx


def pdf_del_articulo(docx):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("Sin LibreOffice: no se hizo el PDF del artículo.")
        return None
    with tempfile.TemporaryDirectory() as perfil:
        subprocess.run(
            [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{perfil}",
             "--convert-to", "pdf", "--outdir", str(BASE), str(docx)],
            check=True, capture_output=True, timeout=600,
        )
    return docx.with_suffix(".pdf")


def main():
    if not ARTICULO.exists():
        sys.exit(f"Falta {ARTICULO}: primero ejecuta python3 articulo/generar_docx.py")
    copia = BASE / "01_articulo_cosecha_de_piedras.docx"
    shutil.copyfile(ARTICULO, copia)
    print(f"Copiado: {copia.name}")
    pdf = pdf_del_articulo(copia)
    if pdf:
        print(f"Escrito: {pdf.name}")
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "referencia.docx"
        plantilla(ref)
        for md in sorted(BASE.glob("0[2-9]_*.md")):
            print(f"Escrito: {convertir(md, ref).name}")


if __name__ == "__main__":
    main()
