"""Genera el .docx del artículo con las normas de Estudios Bolivianos.

Normas aplicadas: tamaño carta, márgenes de 2,5 cm, Times New Roman 12,
interlineado 1,5 (notas al pie en 10 pt, interlineado sencillo).

Uso (desde la raíz del repositorio):
    pip install pypandoc_binary python-docx
    python3 articulo/generar_docx.py

Lee articulo/articulo_piedras_puestas_en_el_agua.md y escribe
articulo/articulo_piedras_puestas_en_el_agua.docx. Al final informa la
extensión en caracteres con espacios (la revista admite 20.000-50.000,
bibliografía incluida) y las palabras de cada resumen (máximo 100).
"""

import re
import subprocess
import zipfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

BASE = Path(__file__).resolve().parent
MD = BASE / "articulo_piedras_puestas_en_el_agua.md"
DOCX = BASE / "articulo_piedras_puestas_en_el_agua.docx"
REF = BASE / "plantilla_estudios_bolivianos.docx"
FONT = "Times New Roman"


def set_font(style, size, bold=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    # Texto en negro (la plantilla de pandoc colorea los títulos).
    color = rpr.find(qn("w:color"))
    if color is not None:
        rpr.remove(color)


def set_spacing(style, line=1.5, before=0, after=6):
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def build_reference():
    data = subprocess.run(
        [pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True,
    ).stdout
    REF.write_bytes(data)
    doc = Document(str(REF))
    styles = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text",
                 "Abstract", "Author", "Date", "Bibliography", "Table", "Definition",
                 "Definition Term", "Image Caption", "Caption", "Table Caption",
                 "Figure", "Captioned Figure"):
        if name in styles:
            set_font(styles[name], 12)
            set_spacing(styles[name])
    for name, size in (("Title", 14), ("Subtitle", 12), ("Heading 1", 12),
                       ("Heading 2", 12), ("Heading 3", 12), ("Heading 4", 12)):
        if name in styles:
            set_font(styles[name], size, bold=True, italic=False)
            set_spacing(styles[name], before=12, after=6)
    for name in ("Image Caption", "Caption", "Table Caption"):
        if name in styles:
            set_font(styles[name], 10, italic=False)
            set_spacing(styles[name], line=1.0, before=3, after=12)
    for name in ("Figure", "Captioned Figure", "Image Caption"):
        if name in styles:
            styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Footnote Text" in styles:
        set_font(styles["Footnote Text"], 10)
        set_spacing(styles["Footnote Text"], line=1.0, after=2)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.5))
    doc.save(str(REF))


def convert():
    pypandoc.convert_file(
        str(MD), "docx", outputfile=str(DOCX),
        extra_args=[f"--reference-doc={REF}", f"--resource-path={BASE}"],
    )
    doc = Document(str(DOCX))
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.5))
    # Tablas en 10 pt con interlineado sencillo para que quepan en la caja.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(2)
                    for run in par.runs:
                        run.font.size = Pt(10)
                        run.font.name = FONT
    doc.save(str(DOCX))


def report():
    doc = Document(str(DOCX))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    with zipfile.ZipFile(DOCX) as z:
        notes_xml = z.read("word/footnotes.xml").decode("utf8") if "word/footnotes.xml" in z.namelist() else ""
    notes = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", notes_xml)
    body = "\n".join(parts)
    total = len(body.replace("\n", " ")) + len(" ".join(notes))
    print(f"Caracteres con espacios (texto + tablas + notas): {total:,}".replace(",", "."))
    md = MD.read_text(encoding="utf8")
    for label in ("Resumen", "Abstract"):
        m = re.search(rf"\*\*{label}\*\*\s*\n\s*\n(.+?)\n", md, re.S)
        if m:
            words = re.sub(r"[*«»\"]", "", m.group(1)).split()
            print(f"Palabras en {label}: {len(words)}")


if __name__ == "__main__":
    build_reference()
    convert()
    report()
    print(f"Escrito: {DOCX}")
